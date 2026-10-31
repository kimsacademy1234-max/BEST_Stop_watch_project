# merge_final_top 블록 다이어그램 (레퍼런스 BEST_stopwatch_presentation_v2.8 5~8p 포맷)
# 사용 : python make_block_diagram.py out.pptx [팀 v2 덱 (v1 → v2 바뀌는 곳 표 1 장을 복사해 옴)]
import copy
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

FONT = "Malgun Gothic"
FONT_TITLE = "Malgun Gothic Semilight"

# 도메인 색 (선, 연한 바탕)
DOM = {
    "com":  ("4A5568", "F3F4F6"),
    "wts":  ("2F855A", "EDF6F0"),
    "sr04": ("2B6CB0", "EBF4FF"),
    "dht":  ("C05621", "FDF0E6"),
    "uart": ("2C7A7B", "E6F4F3"),
    "disp": ("6B46C1", "F3EEFF"),
    "err":  ("C53030", "FDECEC"),
}
PIN = "2D3748"
TXT = "1A202C"
V2G = "1E8E3E"   # v2 : 센서 상태 → 전송 MUX (새 연결)
OFF_LN, OFF_FILL, OFF_TXT, OFF_TAG, OFF_BAND = "D5D9DE", "F7F8F9", "A0A7B0", "E2E5E9", "FAFAFB"


def rgb(h):
    return RGBColor.from_string(h)


class Ctx:
    """한 슬라이드의 상태 : overview 면 도메인 색, 아니면 on / neu / off."""

    def __init__(self, slide, overview, on=(), neu=(), accent="uart", v2=False):
        self.s = slide
        self.overview = overview
        self.v2 = v2
        self.on = set(on)
        self.neu = set(neu)
        self.acc = DOM[accent]
        self.notes = {}

    def state(self, key):
        if self.overview:
            return "dom"
        if key in self.on:
            return "on"
        if key in self.neu:
            return "neu"
        return "off"

    # 상태 + 도메인 → (선 색, 바탕 색, 글자 색, 선 굵기)
    def look(self, key, dom, kind):
        st = self.state(key)
        ln, lt = DOM[dom]
        if st == "dom":
            return dict(ln=ln, fill="FFFFFF" if kind == "block" and dom != "com" else lt, txt=TXT,
                        w=1.5, wire=ln, ww=1.5, tag=(PIN if dom == "com" else ln), tagtxt="FFFFFF", lab=ln, bold=False)
        if st == "on":
            a, al = self.acc
            return dict(ln=a, fill="FFFFFF", txt=TXT, w=2.25, wire=a, ww=2.5, tag=a, tagtxt="FFFFFF", lab=a, bold=True, band=al)
        if st == "neu":
            return dict(ln="4A5568", fill="F3F4F6", txt=TXT, w=1.5, wire="4A5568", ww=1.5, tag=PIN, tagtxt="FFFFFF",
                        lab="4A5568", bold=False, band="F5F6F8")
        return dict(ln=OFF_LN, fill=OFF_FILL, txt=OFF_TXT, w=1.5, wire=OFF_LN, ww=1.0, tag=OFF_TAG, tagtxt=OFF_TXT,
                    lab=OFF_TXT, bold=False, band=OFF_BAND)


# ---------------------------------------------------------------- 기본 도형
def text(slide, x, y, w, h, s, size, color, bold=False, align="l", font=FONT, anchor="m", name="라벨"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = {"m": MSO_ANCHOR.MIDDLE, "t": MSO_ANCHOR.TOP}[anchor]
    lines = s if isinstance(s, list) else [s]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        r = p.add_run()
        r.text = ln
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        _ea(r, font)
    return tb


def text_runs(slide, x, y, w, h, runs, size, align="l", name="라벨"):
    """한 줄에 색 · 굵기가 다른 글자 여러 개 : runs = [(글자, 색, 굵게), ...]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    for s, color, bold in runs:
        r = p.add_run()
        r.text = s
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        _ea(r, FONT)
    return tb


def _ea(run, font):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.5, name="도형"):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.12, 0.08 / min(w, h) * 1.0) if min(w, h) > 0 else 0.1
    return sh


def set_text(sh, lines, color, sizes, bold=True, align="c"):
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (ln, sz) in enumerate(zip(lines, sizes)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[align]
        r = p.add_run()
        r.text = ln
        r.font.name = FONT
        r.font.size = Pt(sz)
        r.font.bold = bold if i == 0 else False
        r.font.color.rgb = rgb(color)
        _ea(r, FONT)


def seg(slide, x1, y1, x2, y2, color, w, arrow=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.name = "선"
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle")
        te.set("w", "med")
        te.set("len", "med")
    return c


# ---------------------------------------------------------------- 다이어그램 요소
def block(cx, key, dom, x, y, w, h, title, sub=None, tsz=12.5):
    L = cx.look(key, dom, "block")
    sh = shape(cx.s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, L["fill"], L["ln"], L["w"], name=f"블록 {title}")
    lines = title.split("\n")
    sizes = [tsz] * len(lines)
    if sub:
        lines.append(sub)
        sizes.append(10)
    set_text(sh, lines, L["txt"], sizes)
    return sh


def band(cx, key, dom, x, y, w, h, label, note=None):
    st = cx.state(key)
    ln, lt = DOM[dom]
    L = cx.look(key, dom, "band")
    fill = lt if st == "dom" else L["band"]
    shape(cx.s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, None, name=f"줄 {label}")
    lab = ln if st == "dom" else (cx.acc[0] if st == "on" else ("4A5568" if st == "neu" else OFF_TXT))
    text(cx.s, x + 0.12, y + 0.03, 6.0, 0.28, label, 12, lab, bold=True)
    if note:
        text(cx.s, x + w - 6.12, y + 0.03, 6.0, 0.28, note, 11, lab, bold=True, align="r")


def tag(cx, key, dom, x, y, w, s, h=0.32, size=11, pin=False):
    L = cx.look(key, dom, "tag")
    st = cx.state(key)
    fill = L["tag"]
    if pin and st in ("dom", "neu"):
        fill = PIN
    sh = shape(cx.s, MSO_SHAPE.PENTAGON, x, y, w, h, fill, None, name=f"태그 {s}")
    sh.adjustments[0] = 0.35 * h / w * 2 if w > 0 else 0.5
    set_text(sh, [s], L["tagtxt"], [size])
    return sh


def wire(cx, key, dom, pts, arrow=True, dash=False, label=None, lpos=None, lw=2.6, lsize=11, lalign="l"):
    L = cx.look(key, dom, "wire")
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        seg(cx.s, x1, y1, x2, y2, L["wire"], L["ww"], arrow=(arrow and i == len(pts) - 2), dash=dash)
    if label:
        lx, ly = lpos
        text(cx.s, lx, ly, lw, 0.26, label, lsize, L["lab"], bold=L["bold"], align=lalign)


def dot(cx, key, dom, x, y):
    L = cx.look(key, dom, "wire")
    r = 0.045
    shape(cx.s, MSO_SHAPE.OVAL, x - r, y - r, 2 * r, 2 * r, L["wire"], None, name="접점")


def gate(cx, key, dom, x, y, sym="&", s=0.6):
    L = cx.look(key, dom, "gate")
    st = cx.state(key)
    if st == "on":
        sh = shape(cx.s, MSO_SHAPE.FLOWCHART_DELAY, x, y, s, s, cx.acc[0], cx.acc[0], 1.75, name="게이트")
        set_text(sh, [sym], "FFFFFF", [12])
    else:
        ln = L["ln"] if st != "dom" else DOM[dom][0]
        sh = shape(cx.s, MSO_SHAPE.FLOWCHART_DELAY, x, y, s, s, "FFFFFF", ln, 1.75, name="게이트")
        set_text(sh, [sym], L["txt"], [12])


def label(cx, key, dom, x, y, w, s, size=11, bold=None, align="l", h=0.26):
    L = cx.look(key, dom, "wire")
    text(cx.s, x, y, w, h, s, size, L["lab"], bold=L["bold"] if bold is None else bold, align=align)


def trap(slide, x, y, w, h, slant, fill, line, lw, name):
    """MUX 기호 : 왼쪽 (입력) 이 넓고 오른쪽 (출력) 이 좁은 사다리꼴."""
    fb = slide.shapes.build_freeform(Inches(x), Inches(y), scale=1.0)
    pts = [(x + w, y + slant), (x + w, y + h - slant), (x, y + h), (x, y)]
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts], close=True)
    sh = fb.convert_to_shape()
    sh.name = name
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(line)
    sh.line.width = Pt(lw)
    return sh


def mux(cx, key, dom, x, y, w, h, title=None, inputs=(), slant=0.25):
    """사다리꼴 MUX · 입력 번호는 왼쪽 안쪽 · 글자 색 (상태별) 을 돌려줌."""
    st = cx.state(key)
    ln, lt = DOM[dom]
    if st == "dom":
        f, l, tc = lt, ln, ln
    elif st == "on":
        f, l, tc = cx.acc[1], cx.acc[0], cx.acc[0]
    elif st == "neu":
        f, l, tc = "F3F4F6", "4A5568", "4A5568"
    else:
        f, l, tc = OFF_FILL, OFF_LN, OFF_TXT
    trap(cx.s, x, y, w, h, slant, f, l, 1.5, name=f"MUX {title}")
    for (iy, t) in inputs:
        text(cx.s, x + 0.06, iy - 0.13, 0.3, 0.26, t, 9.5, tc, bold=True)
    return tc


# ---------------------------------------------------------------- 좌표
ROW_Y = {"wts": 2.72, "sr04": 3.97, "dht": 5.22}  # 줄 (band) 위쪽
ROW_C = {k: v + 0.675 for k, v in ROW_Y.items()}   # 블록 · 게이트 가운데


ROWS = [
    ("wts", "wts", "WTS  스톱워치 · 시계 · 타이머", "wts_top", "control_unit · datapath × 3",
     "disp  시 · 분 · 초 · 1/100", "← ctrl_bit[4:0]  (sw0 · sw1 · sw2 · sw3 · sw4)"),
    ("sr04", "sr04", "SR04  초음파 거리", "sr04_top", "core · mode_ctrl · led_bar",
     "distance_x10 · err · status", "trig → JA1 · echo ← JA2 · sw15 = 거리 막대 LED"),
    ("dht", "dht", "DHT11  온습도", "dht11_top", "core · controller · manager",
     "fnd_data · dashes · err", "dht11_io ↔ JXADC3 (M2)"),
]


def diagram(cx):
    s = cx.s
    # 줄 (band) 는 맨 아래 층
    for k, dom, blabel, *_ in ROWS:
        band(cx, f"band_{k}", dom, 6.35, ROW_Y[k], 6.6, 1.15, blabel, note=cx.notes.get(f"band_{k}"))
    band(cx, "band_uart", "uart", 1.75, 6.95, 16.55, 2.2, "UART  명령 수신 · 응답 전송  (uart_top)",
         note="9600 bps · 8N1")
    # ---------------- 클럭 · 버튼
    tag(cx, "clk", "com", 0.45, 2.30, 1.1, "clk", pin=True)
    wire(cx, "clk", "com", [(1.55, 2.46), (1.9, 2.46)])
    block(cx, "clkgen", "com", 1.9, 2.25, 1.85, 0.42, "clock_tick_generator", tsz=11)
    wire(cx, "clk4k", "com", [(3.75, 2.46), (16.4, 2.46), (16.4, 2.85)], dash=True,
         label="4 kHz  →  디바운서 · FND 스캔", lpos=(4.2, 2.2), lw=3.6)
    wire(cx, "clkgen", "com", [(2.4, 2.67), (2.4, 2.80)])

    for i, b in enumerate("RLUD"):
        y = 2.88 + 0.4 * i
        tag(cx, "btn_pin", "com", 0.45, y, 1.1, f"btn_{b}", pin=True)
        wire(cx, "btn_pin", "com", [(1.55, y + 0.16), (1.9, y + 0.16)])
    block(cx, "deb", "com", 1.9, 2.80, 1.5, 1.62, "btn_debounce\n×4", tsz=12)
    wire(cx, "phy", "com", [(3.4, 3.46), (4.35, 3.46)], label="phy_btn", lpos=(3.48, 3.18), lw=0.9)
    gate(cx, "or", "com", 4.35, 3.31, sym="≥1")
    wire(cx, "orout", "com", [(4.95, 3.61), (5.15, 3.61)], arrow=False)

    # 버튼 버스 (x 5.15) → AND 위 입력
    top, bot = ROW_C["wts"] - 0.15, ROW_C["dht"] - 0.15
    wire(cx, "bus", "com", [(5.15, top), (5.15, bot)], arrow=False)
    dot(cx, "orout", "com", 5.15, 3.61)
    # 선택 버스 (x 5.35) → AND 아래 입력
    wire(cx, "selbus", "com", [(3.95, 5.43), (5.35, 5.43)], arrow=False)
    wire(cx, "selbus", "com", [(5.35, ROW_C["wts"] + 0.15), (5.35, ROW_C["dht"] + 0.15)], arrow=False)
    dot(cx, "selbus", "com", 5.35, 5.43)

    names = {"wts": ("mode_wts", "R·L·U·D"), "sr04": ("mode_dist", "R · U"), "dht": ("mode_dht", "R · U")}
    for k, dom in (("wts", "wts"), ("sr04", "sr04"), ("dht", "dht")):
        c = ROW_C[k]
        wire(cx, f"br_{k}", "com", [(5.15, c - 0.15), (5.55, c - 0.15)])
        if k != "wts":
            dot(cx, "bus", "com", 5.15, c - 0.15)
        wire(cx, f"sel_{k}", "com", [(5.35, c + 0.15), (5.55, c + 0.15)])
        if k != "dht":
            dot(cx, "selbus", "com", 5.35, c + 0.15)
        gate(cx, f"and_{k}", dom, 5.55, c - 0.3)
        label(cx, f"sel_{k}", "com", 5.5, c + 0.31, 1.0, names[k][0], size=10)
        wire(cx, f"a2m_{k}", dom, [(6.15, c), (6.95, c)], label=names[k][1], lpos=(6.2, c - 0.29), lw=0.8, lsize=10.5)

    # ---------------- 모드 선택
    st = cx.state("modebox")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.75, 4.55, 2.3, 1.62, "FBFBFC", "C8CDD3", 1.0, name="모드 선택 칸")
    text(s, 1.82, 4.57, 2.0, 0.28, "모드 선택", 12, "4A5568", bold=True)
    tag(cx, "sw_pin", "com", 0.45, 4.92, 1.1, "sw0 ~ sw6", pin=True, size=10.5)
    wire(cx, "sw_w", "com", [(1.55, 5.08), (2.15, 5.08)])
    tc = mux(cx, "modemux", "com", 2.15, 4.9, 0.5, 1.08, title="모드", inputs=[(5.08, "0"), (5.7, "1")], slant=0.12)
    text(s, 2.65, 4.58, 0.8, 0.26, "2:1 × 7 bit", 9.5, tc, bold=True)
    wire(cx, "mux_out", "com", [(2.65, 5.43), (2.95, 5.43)])
    block(cx, "decode", "com", 2.95, 5.1, 1.0, 0.66, "모드\n디코드", tsz=11)
    label(cx, "decode", "com", 2.55, 5.84, 1.5, "→ mode_* · disp_sel", size=9.5, align="r")

    # ---------------- 모듈 줄
    for k, dom, blabel, title, sub, outlab, note in ROWS:
        y = ROW_Y[k]
        c = ROW_C[k]
        block(cx, k, dom, 6.95, y + 0.3, 2.4, 0.75, title, sub, tsz=13)
        wire(cx, f"o_{k}", dom, [(9.35, c), (13.3, c)], label=outlab, lpos=(9.5, c - 0.29), lw=3.6, lsize=10.5)
        label(cx, f"n_{k}", dom, 9.5, c + 0.05, 3.45, note, size=10)

    # ---------------- FND · LED
    mc = mux(cx, "fndmux", "disp", 13.3, 2.85, 0.7, 3.4, title="FND",
             inputs=[(ROW_C["wts"], "0"), (ROW_C["sr04"], "1"), (ROW_C["dht"], "2")])
    text(s, 13.1, 6.28, 2.2, 0.26, "FND MUX  3:1 × 25 bit", 11, mc, bold=True)
    text(s, 13.1, 6.52, 2.2, 0.26, "sel = disp_sel (2 bit)", 10.5, mc)
    wire(cx, "fmux_out", "disp", [(14.0, 3.3), (15.4, 3.3)], label="fnd_data", lpos=(14.08, 3.03), lw=1.2, lsize=10.5)
    block(cx, "fndctrl", "disp", 15.4, 2.85, 2.0, 1.65, "FND\n컨트롤러", "SW2_FND_CONTROLLER", tsz=14)
    tag(cx, "t_dsel", "disp", 14.4, 3.56, 0.98, "disp_sel", h=0.28, size=10)
    tag(cx, "t_cursor", "wts", 14.4, 3.88, 0.98, "wt_cursor", h=0.28, size=10)
    tag(cx, "t_set2", "wts", 14.4, 4.20, 0.98, "set_sw2", h=0.28, size=10)
    for nm, y in (("fnd_com", 3.18), ("fnd_font", 3.93)):
        wire(cx, "fnd_pins", "disp", [(17.4, y + 0.16), (17.85, y + 0.16)])
        tag(cx, "fnd_pins", "disp", 17.85, y, 1.45, nm)

    block(cx, "led", "disp", 15.4, 4.75, 2.0, 1.5, "LED 로직", "led_output", tsz=14)
    tag(cx, "t_ldsel", "disp", 14.4, 4.84, 0.98, "disp_sel", h=0.28, size=10)
    tag(cx, "t_lwts", "wts", 14.4, 5.16, 0.98, "WTS 신호", h=0.28, size=10)
    tag(cx, "t_lsr", "sr04", 14.4, 5.48, 0.98, "sr04_led", h=0.28, size=10)
    tag(cx, "t_ldht", "dht", 14.4, 5.80, 0.98, "dht_led", h=0.28, size=10)
    wire(cx, "led_pin", "disp", [(17.4, 5.5), (17.85, 5.5)])
    tag(cx, "led_pin", "disp", 17.85, 5.34, 1.45, "led[15:0]")

    # ---------------- UART
    tag(cx, "rx_pin", "uart", 0.45, 7.55, 1.1, "rx", pin=True)
    wire(cx, "rx_w", "uart", [(1.55, 7.71), (1.95, 7.71)])
    block(cx, "urx", "uart", 1.95, 7.4, 1.3, 0.62, "uart_rx", tsz=12)
    wire(cx, "w_rf", "uart", [(3.25, 7.71), (3.55, 7.71)])
    block(cx, "fifo_rx", "uart", 3.55, 7.4, 1.0, 0.62, "fifo", "RX × 8", tsz=12)
    wire(cx, "w_fd", "uart", [(4.55, 7.71), (4.85, 7.71)])
    block(cx, "dec", "uart", 4.85, 7.3, 2.0, 0.82, "ASCII_DECODER", "/명령 + 옵션 파싱", tsz=12)
    wire(cx, "w_dc", "uart", [(6.85, 7.71), (7.4, 7.71)])
    block(cx, "baud", "uart", 1.95, 8.4, 1.3, 0.58, "baud_tick_gen", "9600 × 16", tsz=10.5)
    wire(cx, "baud", "uart", [(2.6, 8.4), (2.6, 8.02)])
    label(cx, "baud", "uart", 3.35, 8.56, 2.0, "→ RX · TX 공용", size=10)

    block(cx, "ctrl", "uart", 7.4, 7.3, 2.6, 1.5, "UART_CONTROLLER", "mode · auto · btn FSM", tsz=12.5)
    phy_tag(cx)
    # 컨트롤러 → 위쪽 (모드 · 버튼)
    wire(cx, "w_cbit", "uart", [(7.55, 7.3), (7.55, 6.8), (1.85, 6.8), (1.85, 5.7), (2.15, 5.7)],
         label="module_control_bit[6:0]", lpos=(4.3, 6.82), lw=2.6, lsize=10)
    wire(cx, "w_umode", "uart", [(7.85, 7.3), (7.85, 6.65), (2.4, 6.65), (2.4, 5.92)],
         label="uart_mode", lpos=(2.5, 6.40), lw=1.2, lsize=10)
    wire(cx, "w_btn", "uart", [(8.15, 7.3), (8.15, 6.5), (4.1, 6.5), (4.1, 3.76), (4.35, 3.76)],
         label="uart_btn", lpos=(4.2, 4.05), lw=0.9, lsize=10)

    # 1 Hz tick (auto)
    block(cx, "tick", "uart", 10.45, 7.62, 1.5, 0.62, "sender_tick", "1 Hz · scan 0→4", tsz=11)
    wire(cx, "tick", "uart", [(10.45, 7.93), (10.0, 7.93)])
    # send_go
    wire(cx, "w_go", "uart", [(10.0, 7.42), (15.17, 7.42), (15.17, 7.75)],
         label="send_go", lpos=(11.2, 7.16), lw=1.0, lsize=10)
    # 전송 MUX : 8:1 × 24 bit (send_sel 3 bit) · 입력은 실제 출처에서 선으로
    tc = mux(cx, "smux", "uart", SMX, SMY, SMW, SMH, title="전송", slant=SMS)
    text(s, 12.95, 7.14, 2.4, 0.26, "전송 MUX  8:1 × 24 bit", 10.5, tc, bold=True)
    send_inputs(cx, tc)
    wire(cx, "w_sm", "uart", [(SMX + SMW, 8.2), (14.45, 8.2)])
    # send_sel
    wire(cx, "w_sel", "uart", [(9.7, 8.8), (9.7, 9.02), (15.17, 9.02), (15.17, 8.65)],
         label="send_sel", lpos=(10.4, 8.76), lw=1.0, lsize=10)
    wire(cx, "w_selmux", "uart", [(13.65, 9.02), (13.65, SMY + SMH - SMS * (13.65 - SMX) / SMW)])
    dot(cx, "w_sel", "uart", 13.65, 9.02)

    block(cx, "sender", "uart", 14.45, 7.75, 1.45, 0.9, "uart_sender", "숫자 → ASCII 한 줄", tsz=11.5)
    wire(cx, "w_sf", "uart", [(15.9, 8.2), (16.15, 8.2)])
    block(cx, "fifo_tx", "uart", 16.15, 7.89, 0.9, 0.62, "fifo", "TX × 64", tsz=11.5)
    wire(cx, "w_ft", "uart", [(17.05, 8.2), (17.3, 8.2)])
    block(cx, "utx", "uart", 17.3, 7.89, 0.95, 0.62, "uart_tx", tsz=11.5)
    wire(cx, "w_tx", "uart", [(18.25, 8.2), (18.45, 8.2)])
    tag(cx, "tx_pin", "uart", 18.45, 8.04, 0.9, "tx")
    if cx.v2 and cx.state("v2") != "off":
        for y in (4.0, 5.25):   # 센서 상태가 FND 와 전송 둘 다로
            text_runs(s, 9.15, y, 3.55, 0.28, [("상태 → FND MUX ", V2G, True), ("+ 전송 MUX", V2G, True)],
                      10, align="r", name="라벨 v2")


# 전송 MUX 자리 (왼쪽 x · 위 y · 폭 · 왼쪽 높이 · 기울기)
SMX, SMY, SMW, SMH, SMS = 13.0, 7.55, 0.95, 1.4, 0.18


def send_inputs(cx, tc):
    """전송 MUX 입력 : 0 · 1 · 2 = wts_top 시간 3 개, 3 = sr04_top, 4 = dht11_top, 5~7 = 24'd0 (7 = ERROR).
    센서마다 선 1 개 = MUX 입력 1 개 (v2 는 값 + 상태 비트를 한 묶음으로).
    세로선은 아래 입력일수록 왼쪽 → 서로 안 겹침."""
    s = cx.s
    xw, xs, xd = 12.88, 12.78, 12.5
    # WTS : 시간 3 개 → 0 · 1 · 2
    wire(cx, "st_wts", "wts", [(xw, ROW_C["wts"]), (xw, 8.00), (SMX, 8.00)])
    for y in (7.68, 7.84):
        wire(cx, "st_wts", "wts", [(xw, y), (SMX, y)])
        dot(cx, "st_wts", "wts", xw, y)
    dot(cx, "st_wts", "wts", xw, ROW_C["wts"])
    wire(cx, "st_dist", "sr04", [(xs, ROW_C["sr04"]), (xs, 8.22), (SMX, 8.22)])
    dot(cx, "st_dist", "sr04", xs, ROW_C["sr04"])
    wire(cx, "st_dht", "dht", [(xd, ROW_C["dht"]), (xd, 8.44), (SMX, 8.44)])
    dot(cx, "st_dht", "dht", xd, ROW_C["dht"])
    # 핀 번호 · 이름 (MUX 안쪽)
    g = V2G if cx.state("smux") != "off" else tc
    pins = [(7.68, [("0 sw", tc)]), (7.84, [("1 tm", tc)]), (8.00, [("2 wt", tc)]),
            (8.22, [("3 거리", tc)] + ([("·상태", g)] if cx.v2 else [])),
            (8.44, [("4 온습도", tc)] + ([("·상태", g)] if cx.v2 else [])),
            (8.72, [("5~7 = 0", tc)])]
    for y, runs in pins:
        text_runs(s, SMX + 0.05, y - 0.1, 0.88, 0.2, [(t, c, True) for t, c in runs], 8.5, name="MUX 핀")


def phy_tag(cx):
    """디바운서 출력 phy_btn → UART_CONTROLLER (i_btn_R/L/U/D).
    같은 clk 에 물리 버튼이 있으면 컨트롤러가 uart_btn 을 막고 저장 → 다음 clk 에 출력 (물리 먼저)."""
    tag(cx, "t_phy", "com", 6.42, 8.2, 0.98, "phy_btn", h=0.28, size=10)
    label(cx, "t_phy", "com", 4.32, 8.21, 2.0, "같은 clk 면 물리 먼저", size=9.5, align="r")


def v2_note(cx, head, body):
    text_runs(cx.s, 0.92, 10.48, 16.8, 0.32, [(head, V2G, True), ("  " + body, "4A5568", False)], 11,
              name="라벨 v2")


# ---------------------------------------------------------------- 슬라이드 틀
HEADER = "WTS · SR04 · DHT11 · UART"


def frame(prs, section, title, summary, chip_color, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 13.6, 0.22, 6.0, 0.3, HEADER, 10.5, "8C8C8C", align="r")
    shape(s, MSO_SHAPE.OVAL, 0.92, 0.55, 0.14, 0.14, "404040", None, name="shape")
    text(s, 1.14, 0.44, 10.0, 0.34, section, 12.5, "404040", bold=True)
    text(s, 0.92, 0.78, 18.2, 0.72, title, 30, "262626", font=FONT_TITLE)
    chip = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.92, 1.66, 0.95, 0.44, chip_color, None, name="chip")
    chip.adjustments[0] = 0.5
    set_text(chip, ["요약"], "FFFFFF", [14])
    text(s, 2.02, 1.62, 17.1, 0.52, summary, 19, "262626", bold=True)
    text(s, 18.2, 10.62, 0.9, 0.35, f"{page:02d}", 11, "8C8C8C", align="r")
    return s


def badges(cx, pts):
    a = cx.acc[0]
    for n, (x, y) in enumerate(pts, 1):
        b = shape(cx.s, MSO_SHAPE.OVAL, x, y, 0.38, 0.38, a, "FFFFFF", 1.5, name=f"단계 {n}")
        set_text(b, [str(n)], "FFFFFF", [13])


def steps(cx, items):
    a = cx.acc[0]
    for i, t in enumerate(items):
        col, row = i % 3, i // 3
        x = 0.92 + 6.1 * col
        y = 9.42 + 0.5 * row
        b = shape(cx.s, MSO_SHAPE.OVAL, x, y, 0.36, 0.36, a, None, name=f"단계표 {i + 1}")
        set_text(b, [str(i + 1)], "FFFFFF", [13])
        if isinstance(t, list):   # [(글자, 색), ...]
            text_runs(cx.s, x + 0.45, y - 0.01, 5.65, 0.38, [(a_, c_, True) for a_, c_ in t], 13.5)
        else:
            text(cx.s, x + 0.45, y - 0.01, 5.6, 0.38, t, 13.5, "262626", bold=True)


def legend(cx):
    s = cx.s
    x = 0.92
    for dom, nm in (("wts", "WTS"), ("sr04", "SR04 거리"), ("dht", "DHT11 온습도"), ("uart", "UART"), ("disp", "FND · LED")):
        ln, lt = DOM[dom]
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 9.62, 0.34, 0.24, lt, ln, 1.5, name="범례")
        text(s, x + 0.42, 9.56, 1.7, 0.36, nm, 13, "262626", bold=True)
        x += 0.42 + (len(nm) * 0.13 + 0.55)
    x += 0.2
    g = shape(s, MSO_SHAPE.FLOWCHART_DELAY, x, 9.55, 0.4, 0.38, "FFFFFF", "4A5568", 1.5, name="범례 AND")
    set_text(g, ["&"], TXT, [10])
    text(s, x + 0.5, 9.56, 2.6, 0.36, "AND · ≥1 = OR 게이트", 13, "262626", bold=True)
    x += 3.25
    seg(s, x, 9.74, x + 0.6, 9.74, "4A5568", 1.5, arrow=True, dash=True)
    text(s, x + 0.7, 9.56, 1.5, 0.36, "4 kHz 클럭", 13, "262626", bold=True)
    x += 2.25
    t = shape(s, MSO_SHAPE.PENTAGON, x, 9.58, 0.55, 0.3, PIN, None, name="범례 태그")
    text(s, x + 0.65, 9.56, 4.0, 0.36, "태그 = 보드 핀 · 같은 이름의 신호", 13, "262626", bold=True)
    text(s, 0.92, 10.05, 17.0, 0.34,
         "명령 형식 : /<명령> <옵션>  ·  /mod  /btn  /get  /auto  /set  /off  →  다음 장부터 명령별 흐름 (켜진 경로만 색, 나머지 회색)",
         12, "4A5568", bold=True)


# ---------------------------------------------------------------- 슬라이드별 내용
RX = ["rx_pin", "rx_w", "urx", "w_rf", "fifo_rx", "w_fd", "dec", "w_dc", "ctrl", "baud", "band_uart"]
TX = ["w_go", "w_sel", "sender", "w_sf", "fifo_tx", "w_ft", "utx", "w_tx", "tx_pin"]
CLK = ["clk", "clkgen", "clk4k"]
FND = ["fndmux", "fmux_out", "fndctrl", "t_dsel", "fnd_pins", "led", "t_ldsel", "led_pin"]
MODS = ["band_wts", "wts", "o_wts", "n_wts", "band_sr04", "sr04", "o_sr04", "n_sr04",
        "band_dht", "dht", "o_dht", "n_dht"]
MODESEL = ["modemux", "mux_out", "decode", "selbus"]


def btn_slide(k, dom):
    on = RX + ["w_btn", "or", "orout", "bus", f"br_{k}", f"sel_{k}", f"and_{k}", f"a2m_{k}",
               f"band_{k}", k, f"o_{k}", f"n_{k}", "t_phy"] + FND
    on += {"wts": ["t_lwts"], "sr04": ["t_lsr"], "dht": ["t_ldht"]}[k]
    if k == "wts":
        on += ["t_cursor"]
    neu = CLK + ["btn_pin", "deb", "phy"] + MODESEL + ["w_cbit", "w_umode", "modemux", "sw_pin", "sw_w"]
    return on, neu


SLIDES = [
    dict(overview=True, accent="err", chip="C00000", section="02  시스템 구조",
         title="Top 블록 다이어그램 — UART 명령에서 FND · LED · TX 까지",
         summary="명령은 RX → 디코더 → 컨트롤러에서 갈라짐 · 모드 비트 / 버튼 / 전송 요청으로 각 모듈 · 화면 · TX 까지"),

    dict(accent="disp", v2=True, section="02  시스템 구조 · 보드 입력",
         title="스위치 · 버튼 흐름 — /mod 전에는 보드 입력으로 조작",
         summary="리셋 후 첫 /mod 전까지 sw0 ~ sw6 이 모드 결정 · 버튼 4 개는 디바운스 → OR → 지금 화면 모듈의 AND 만 통과",
         on=CLK + ["btn_pin", "deb", "phy", "or", "orout", "bus", "br_wts", "br_sr04", "br_dht",
                   "sel_wts", "sel_sr04", "sel_dht", "and_wts", "and_sr04", "and_dht",
                   "a2m_wts", "a2m_sr04", "a2m_dht", "sw_pin", "sw_w", "modemux", "mux_out", "decode", "selbus"]
            + MODS + FND + ["t_cursor", "t_set2", "t_lwts", "t_lsr", "t_ldht"],
         neu=["t_phy"],
         ann=[(0.45, 5.28, 1.25, "첫 /mod 전")],
         badges=[(1.22, 4.5), (3.5, 4.68), (3.45, 2.72), (4.05, 2.95), (5.45, 2.68), (14.2, 2.6)],
         steps=["sw0 ~ sw6 → 모드 선택 MUX 0 번 (uart_mode = 0)",
                "모드 디코드 : sw6 온습도 > sw5 거리 > WTS",
                "btn R · L · U · D → 디바운서 (4 kHz × 8 샘플)",
                "OR : 물리 btn | uart_btn → 같은 버튼 tick",
                "AND : 지금 화면 mode_* 모듈만 통과",
                "disp_sel → FND MUX · LED (sw15 거리 막대)"],
         v2_note=("참고", "sw0 ~ sw4 → WTS 기능 (시계 · 스톱워치 · 설정 · 12/24 · 타이머)  ·  "
                          "첫 /mod 를 받으면 스위치 무시 · btnC 리셋하면 다시 스위치  →  다음 장 /mod 흐름")),

    dict(accent="uart", v2=True, section="02  시스템 구조 · 명령별 흐름",
         title="/mod 흐름 — 화면 · 버튼 대상 바꾸기",
         summary="첫 /mod 부터 스위치 대신 UART 비트가 모드 결정 · [6:5] 로 버튼 대상과 FND · LED 화면이 같이 바뀜",
         on=RX + ["w_cbit", "w_umode", "modemux", "mux_out", "decode", "selbus", "sel_wts", "sel_sr04", "sel_dht",
                  "n_wts"] + FND + ["t_ldsel"],
         neu=CLK + MODS + ["deb", "btn_pin", "phy", "or", "orout", "bus", "br_wts", "br_sr04", "br_dht",
                           "and_wts", "and_sr04", "and_dht", "a2m_wts", "a2m_sr04", "a2m_dht",
                           "t_cursor", "t_set2", "t_lwts", "t_lsr", "t_ldht"],
         v2_note=("v2", "이 흐름은 v1 과 같음 · 바뀐 곳은 전송 MUX 입력 3 · 4 (거리·상태 · 온습도·상태) 뿐  →  /get · /auto (v2) 장 참고"),
         ann=[(0.45, 5.28, 1.25, "/mod 후 무시")],
         badges=[(6.66, 7.08), (7.0, 8.55), (1.3, 6.05), (3.5, 4.68), (9.1, 3.55), (14.2, 2.6)],
         steps=["rx → uart_rx → FIFO → 디코더 : \"/mod \" + 4 글자",
                "컨트롤러 : w_hm · s_ms · t_hm · dist · temp → 7 비트",
                "uart_mode = 1 (리셋 전까지) → MUX 가 UART 값 사용",
                "모드 디코드 [6:5] → mode_wts · dist · dht 버튼 대상",
                "ctrl_bit[4:0] → WTS 스위치 기능 (bit3 12/24 는 유지)",
                "disp_sel → FND MUX · LED 가 그 화면 값 선택"]),

    dict(accent="wts", section="02  시스템 구조 · 명령별 흐름",
         title="/btn 흐름 ① WTS 화면 — /btn R · L · U · D",
         summary="UART 버튼은 물리 버튼과 OR → mode_wts AND 만 통과 · 같은 clk 에 겹치면 물리 버튼 먼저",
         btn="wts",
         v2=True,
         v2_note=("v2", "이 흐름은 v1 과 같음 · 바뀐 곳은 전송 MUX 입력 3 · 4 (거리·상태 · 온습도·상태) 뿐  →  /get · /auto (v2) 장 참고"),
         badges=[(6.66, 7.08), (8.3, 6.12), (4.05, 2.95), (5.45, 2.68), (9.1, 2.95), (14.2, 2.6)],
         steps=["디코더 : \"/btn \" + 1 글자 (R · L · U · D)",
                "컨트롤러 : uart_btn 1 clk 펄스 (물리 버튼 clk 엔 대기)",
                "OR : 물리 btn | uart_btn → 같은 버튼 tick",
                "AND : mode_wts = 1 → WTS 만 통과 (거리 · 온습도 차단)",
                "wts_top : 스톱워치 R 시작 · 시계 커서 · 타이머 +1분",
                "FND MUX 0 번 (WTS) → FND 표시 · LED"]),

    dict(accent="sr04", section="02  시스템 구조 · 명령별 흐름",
         title="/btn 흐름 ② 거리 화면 — /btn R 측정 · /btn U 자동",
         summary="/mod dist 상태에서 R · U 만 SR04 로 · trig 펄스 → echo 폭 측정 → 0.1 cm 표시 · 막대 LED",
         btn="sr04",
         v2=True,
         v2_note=("v2", "이 흐름은 v1 과 같음 · 바뀐 곳은 전송 MUX 입력 3 · 4 (거리·상태 · 온습도·상태) 뿐  →  /get · /auto (v2) 장 참고"),
         badges=[(6.66, 7.08), (8.3, 6.12), (4.05, 2.95), (5.45, 3.93), (9.1, 4.2), (14.2, 2.6)],
         steps=["디코더 : \"/btn R\" 또는 \"/btn U\"",
                "컨트롤러 : uart_btn 1 clk 펄스 (물리 버튼 clk 엔 대기)",
                "OR : 물리 btn | uart_btn → 같은 버튼 tick",
                "AND : mode_dist = 1 → SR04 만 통과",
                "sr04_top : R 1 회 측정 · U 자동 측정 (60 ms) 켜기/끄기",
                "FND MUX 1 번 → \"352.7\" / Err1~3 · LED 거리 막대"]),

    dict(accent="dht", section="02  시스템 구조 · 명령별 흐름",
         title="/btn 흐름 ③ 온습도 화면 — /btn R 측정 · /btn U 자동",
         summary="/mod temp 상태에서 R · U 만 DHT11 로 · 단선 dht11_io 로 40 비트 수신 → 습도.온도 표시",
         btn="dht",
         v2=True,
         v2_note=("v2", "이 흐름은 v1 과 같음 · 바뀐 곳은 전송 MUX 입력 3 · 4 (거리·상태 · 온습도·상태) 뿐  →  /get · /auto (v2) 장 참고"),
         badges=[(6.66, 7.08), (8.3, 6.12), (4.05, 2.95), (5.45, 5.18), (9.1, 5.45), (14.2, 2.6)],
         steps=["디코더 : \"/btn R\" 또는 \"/btn U\"",
                "컨트롤러 : uart_btn 1 clk 펄스 (물리 버튼 clk 엔 대기)",
                "OR : 물리 btn | uart_btn → 같은 버튼 tick",
                "AND : mode_dht = 1 → DHT11 만 통과",
                "dht11_top : R 1 회 측정 · U 자동 측정 (2 s) 켜기/끄기",
                "FND MUX 2 번 → \"55.24\" / ---- / Err · LED 상태"]),

    dict(accent="uart", section="02  시스템 구조 · 명령별 흐름",
         title="/get 흐름 — 지금 화면 값 1 회 전송",
         summary="컨트롤러가 현재 모드로 send_sel 을 정하고 send_go 1 펄스 · sender 가 그 순간 값을 ASCII 한 줄로 TX",
         on=RX + TX + ["smux", "st_wts", "st_dist", "st_dht", "w_sm", "w_selmux"],
         neu=CLK + MODS + FND + ["t_cursor", "t_set2", "t_lwts", "t_lsr", "t_ldht"],
         badges=[(6.66, 7.08), (9.3, 9.0), (12.06, 8.56), (15.28, 7.0), (14.8, 8.62), (18.3, 7.55)],
         steps=["디코더 : \"/get \" + Enter",
                "컨트롤러 : /mod 모드 → send_sel (온습도 4 · 거리 3 · …)",
                "전송 MUX : *_time · distance · hum·temp 중 선택",
                "send_go 1 펄스 → sender 가 그 순간 값 캡처",
                "sender_char_mux : \"14:05:09.12\" · \"045cm\" · \"55%24C\"",
                "FIFO_TX → uart_tx → tx · /auto 중엔 /get 무시"]),

    dict(accent="uart", section="02  시스템 구조 · 명령별 흐름",
         title="/auto 흐름 — 1 초마다 자동 전송",
         summary="컨트롤러가 AUTO 상태로 1 Hz tick 을 켬 · 매 tick 마다 정해진 (all 이면 0→4 순환) 데이터를 한 줄씩 TX",
         on=RX + TX + ["smux", "st_wts", "st_dist", "st_dht", "w_sm", "w_selmux", "tick"],
         neu=CLK + MODS + FND + ["t_cursor", "t_set2", "t_lwts", "t_lsr", "t_ldht"],
         badges=[(6.66, 7.08), (7.0, 8.55), (10.55, 8.32), (12.06, 8.56), (14.8, 8.62), (18.3, 7.55)],
         steps=["디코더 : \"/auto \" + 3 글자 (_sw _tm _wh dis tem all off)",
                "컨트롤러 : AUTO 상태 · sender_tick 켜기",
                "1 Hz 마다 send_sel · all = 0→1→2→3→4 순환",
                "전송 MUX → sender (busy 면 저장 후 바로 이어서)",
                "FIFO_TX → uart_tx → tx 한 줄씩",
                "/auto off → 멈춤 · FND 화면은 그대로 (/mod 와 별개)"]),

    dict(accent="wts", section="02  시스템 구조 · 명령별 흐름",
         title="/set · /off 흐름 — 시계 시간 설정 · 12 / 24 시간",
         summary="/set 은 모드 비트 bit2 (set_sw2) · bit3 (hour24_sw3) 만 바꿈 · 값 변경은 /btn 으로 · /mod 받은 뒤에만 동작",
         on=RX + ["w_cbit", "w_umode", "modemux", "mux_out", "decode", "n_wts", "band_wts", "wts", "o_wts",
                  "fndmux", "fmux_out", "fndctrl", "t_cursor", "t_set2", "fnd_pins", "led", "t_lwts", "led_pin"],
         neu=CLK + ["deb", "btn_pin", "phy", "or", "orout", "bus", "br_wts", "and_wts", "a2m_wts", "sel_wts",
                    "selbus", "w_btn", "t_dsel", "t_ldsel"],
         badges=[(6.66, 7.08), (7.0, 8.55), (1.3, 6.05), (9.1, 3.55), (5.45, 2.68), (14.2, 2.6)],
         steps=["디코더 : \"/set w_hm\" · \"/set 24h⏎\" · \"/off w_h\"",
                "/set w_hm → bit2 = 1 (w_hm 화면일 때만) · /off → 0",
                "/set 24h · 12h → bit3 · /mod 를 바꿔도 유지",
                "ctrl_bit → wts_top 시계 설정 모드 · 12/24 변환",
                "/btn L · R 커서 · U · D ± (버튼 흐름은 /btn ① 과 같음)",
                "FND : 커서 자리 0.5 s 점멸 · LD15 오후"]),

    dict(accent="err", section="02  시스템 구조 · 명령별 흐름",
         title="오류 흐름 — 형식이 틀리면 ERROR 응답",
         summary="디코더가 틀린 글자에서 바로 버리고 o_error · 컨트롤러는 ERROR 를 get / auto 보다 먼저 보냄",
         on=RX + TX,
         neu=CLK + MODS + FND + ["t_cursor", "t_set2", "t_lwts", "t_lsr", "t_ldht"],
         badges=[(6.66, 7.08), (7.0, 8.55), (9.3, 9.0), (14.8, 8.62), (18.3, 7.55), (3.5, 4.68)],
         steps=["디코더 : 모르는 명령 · 공백 없음 · 빈 옵션 → 버림",
                "o_error 1 clk → 컨트롤러 err_pend 저장",
                "ERROR > get / auto · sender busy 면 대기",
                "send_sel = 7 → sender 가 \"ERROR\" 글자 생성",
                "FIFO_TX → uart_tx → tx : \"ERROR\" CR LF",
                "모드 · 버튼 · 화면 그대로 · 다음 \"/\" 부터 다시"]),

    # ---------------- v2 (팀 덱 FINAL/block_diagram 10 · 11 · 16 · 17 장 문구)
    dict(copy=10),   # v1 → v2 바뀌는 곳 (표 · 그림 없음) : 팀 덱에서 그대로

    dict(overview=True, v2=True, accent="err", chip="C00000", section="02  시스템 구조  ·  v2",
         title="Top 블록 다이어그램 — UART 명령에서 FND · LED · TX 까지  (v2)",
         summary="명령은 RX → 디코더 → 컨트롤러에서 갈라짐 · 모드 비트 / 버튼 / 전송 요청으로 각 모듈 · 화면 · TX 까지",
         v2_note=("v2", "센서 상태 비트를 전송 MUX 입력 3 · 4 에 같이 실음  (3 = {13'd0, 상태, 거리} · "
                        "4 = {4'd0, err, dashes, 상태, 습도, 온도})  ·  블록 · 포트 · 24 bit 폭은 그대로")),

    dict(accent="uart", v2=True, section="02  시스템 구조 · 명령별 흐름  ·  v2",
         title="/get 흐름 — 지금 화면 값 1 회 전송  (v2)",
         summary="전송 MUX 에 상태 비트를 같이 실음 (24 비트의 빈 윗비트) → sender 가 상태를 보고 Err N · --%--C 글자",
         on=RX + TX + ["smux", "st_wts", "st_dist", "st_dht", "w_sm", "w_selmux", "v2"],
         neu=CLK + MODS + FND + ["t_cursor", "t_set2", "t_lwts", "t_lsr", "t_ldht"],
         badges=[(6.66, 7.08), (9.3, 9.0), (12.06, 8.56), (15.28, 7.0), (14.8, 8.62), (18.3, 7.55)],
         steps=["디코더 : \"/get \" + Enter",
                "컨트롤러 : /mod 모드 → send_sel (온습도 4 · 거리 3 · …)",
                "전송 MUX : *_time · 거리·상태 · 온습도·상태",
                "send_go 1 펄스 → sender 가 그 순간 값 캡처",
                [("sender : \"045cm\" · \"55%24C\"", "262626"), (" + \"Err 2\" · \"--%--C\"", V2G)],
                "FIFO_TX → uart_tx → tx · /auto 중엔 /get 무시"],
         v2_note=("v2", "전송 MUX 3 = {13'd0, 상태, 거리} · 4 = {4'd0, err, dashes, 상태, 습도, 온도}  →  "
                        "PC 도 Err 1 · Err 2 · Err 3 · --%--C")),

    dict(accent="uart", v2=True, section="02  시스템 구조 · 명령별 흐름  ·  v2",
         title="/auto 흐름 — 1 초마다 자동 전송  (v2)",
         summary="컨트롤러가 AUTO 상태로 1 Hz tick 을 켬 · 매 tick 마다 정해진 (all 이면 0→4 순환) 데이터를 한 줄씩 TX",
         on=RX + TX + ["smux", "st_wts", "st_dist", "st_dht", "w_sm", "w_selmux", "tick", "v2"],
         neu=CLK + MODS + FND + ["t_cursor", "t_set2", "t_lwts", "t_lsr", "t_ldht"],
         badges=[(6.66, 7.08), (7.0, 8.55), (10.55, 8.32), (12.06, 8.56), (14.8, 8.62), (18.3, 7.55)],
         steps=["디코더 : \"/auto \" + 3 글자 (_sw _tm _wh dis tem all off)",
                "컨트롤러 : AUTO 상태 · sender_tick 켜기",
                "1 Hz 마다 send_sel · all = 0→1→2→3→4 순환",
                "전송 MUX (상태 비트 포함) → sender (busy 면 저장 후 이어서)",
                "FIFO_TX → uart_tx → tx 한 줄씩",
                "/auto off → 멈춤 · FND 화면은 그대로 (/mod 와 별개)"],
         v2_note=("v2", "/auto dis · tem · all 의 거리 · 온습도 줄도 /get 과 같은 길  →  "
                        "오류면 Err N · 측정 전이면 --%--C")),
]


def copy_slide(prs, src, n, page):
    """다른 덱 (src) 의 n 장을 도형 그대로 복사 (그림 · 차트 없는 장 전용) · 쪽 번호만 새로."""
    ss = src.slides[n - 1]
    d = prs.slides.add_slide(prs.slide_layouts[6])
    tree = d.shapes._spTree
    for el in list(tree)[2:]:          # nvGrpSpPr · grpSpPr 만 남김
        tree.remove(el)
    for el in list(ss.shapes._spTree)[2:]:
        tree.append(copy.deepcopy(el))
    hit = [sh for sh in d.shapes if sh.has_text_frame and sh.text_frame.text == str(n)
           and abs(sh.left - Inches(18.2)) < Inches(0.05)]
    assert len(hit) == 1, (n, len(hit))
    hit[0].text_frame.paragraphs[0].runs[0].text = f"{page:02d}"


def build(out, v2_src=None):
    prs = Presentation()
    prs.slide_width = Inches(20)
    prs.slide_height = Inches(11.25)
    src = Presentation(v2_src) if v2_src else None
    page = 0
    for d in SLIDES:
        if "copy" in d:
            if src:
                page += 1
                copy_slide(prs, src, d["copy"], page)
            continue
        page += 1
        acc = DOM[d["accent"]][0]
        s = frame(prs, d["section"], d["title"], d["summary"], d.get("chip", acc), page)
        if d.get("btn"):
            on, neu = btn_slide(d["btn"], d["accent"])
        else:
            on, neu = d.get("on", ()), d.get("neu", ())
        cx = Ctx(s, d.get("overview", False), on, neu, d["accent"], v2=d.get("v2", False))
        cx.notes = d.get("notes", {})
        if d.get("btn"):
            for k in ("wts", "sr04", "dht"):
                if k != d["btn"]:
                    cx.notes[f"band_{k}"] = "버튼 차단 (mode = 0) · 진행 중 동작은 계속"
        diagram(cx)
        for (x, y, w, t) in d.get("ann", ()):
            text(s, x, y, w, 0.28, t, 11, cx.acc[0], bold=True)
        if cx.overview:
            legend(cx)
        else:
            badges(cx, d["badges"])
            steps(cx, d["steps"])
        if d.get("v2_note"):
            v2_note(cx, *d["v2_note"])
    prs.save(out)


if __name__ == "__main__":
    # python make_block_diagram.py out.pptx [FINAL/block_diagram/merge_final_block_diagram.pptx]
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
