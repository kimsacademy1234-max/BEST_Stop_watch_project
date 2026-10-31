# -*- coding: utf-8 -*-
"""'Basys3 보드 IO 구성' 슬라이드 본문을 핀 연결도 형식으로 다시 그림 (PPT 도형).

원본 형식 : FINAL/docs/basys3_pinmap.png
- 프레임 (머리글 · 섹션 · 제목 · 요약 · 쪽 번호) 과 보드 사진은 그대로 두고 나머지 도형을 지운 뒤 다시 그림
- 다시 실행해도 같은 결과 (idempotent)

    python3 board_io_slide.py      (FINAL_presentation.pptx 제자리 수정)
"""
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "FINAL_presentation.pptx")
TITLE = "Basys3 보드 IO 구성"

# ---------------------------------------------------------------- 색 (덱 팔레트)
GRAY = "595959"     # GND · VCC
GREEN = "2F855A"    # SR04 ECHO
RED = "C00000"      # SR04 TRIG · BTNC(리셋) · 전원 주의
ORANGE = "C05621"   # DHT11 DATA · sw15
UART = "6B46C1"     # micro USB (UART)
BLUE = "2B6CB0"     # sw0 ~ sw4 (WTS)
TEAL = "0F766E"     # sw5 · sw6 · 버튼
TEXT = "262626"
SUB = "595959"
FONT = "Malgun Gothic"

# 파란 보드 사진 위 테두리는 진한 색이 묻혀서 같은 계열 밝은 색 사용
ON_BOARD = {GRAY: "CBD5E0", GREEN: "68D391", RED: "FC8181", ORANGE: "F6AD55",
            UART: "B794F4", BLUE: "90CDF4", TEAL: "4FD1C5"}

# ---------------------------------------------------------------- 배치 (in)
PIC_X, PIC_Y, PIC_W = 3.75, 5.64, 8.52


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- 도형 도우미
def style_run(run, size, color, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text_box(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
             shape=None, fill=None, line=None, line_w=1.25, margin=0.0):
    """lines = [(text, size, color, bold), ...]  한 항목 = 한 문단.
    text 에 [(text, size, color, bold), ...] 리스트를 주면 한 문단 안 여러 run."""
    if shape is None:
        sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = 0.12
        sh.shadow.inherit = False
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    elif shape is not None:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_w)
    elif shape is not None:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin if side in ("left", "right") else 0))
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = item if isinstance(item, list) else [item]
        for text, size, color, bold in runs:
            style_run(p.add_run(), size, color, bold)
            p.runs[-1].text = text
    return sh


def box(slide, x, y, w, h, color, fill=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, width=2.25):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.15
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(width)
    return sh


def wire(slide, x1, y1, x2, y2, color, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def P(rx, ry, pic_h):
    """보드 사진 상대 좌표 -> 슬라이드 좌표."""
    return PIC_X + rx * PIC_W, PIC_Y + ry * pic_h


# ---------------------------------------------------------------- 본문
def build(slide):
    pic = next(sh for sh in slide.shapes if sh.shape_type == 13)
    pic_h = PIC_W * pic.height / pic.width
    pic.left, pic.top, pic.width, pic.height = (Inches(PIC_X), Inches(PIC_Y),
                                                Inches(PIC_W), Inches(pic_h))

    # ---- 왼쪽 : Pmod 핀 배열 (JA · JXADC) ----------------------------------
    pitch, cell = 0.235, 0.17
    lab_l, cell_x1, cell_x2, lab_r, lab_w = 0.85, 2.32, 2.54, 2.78, 1.42
    used = {}   # 라벨 -> (오른쪽 끝 x, 중심 y)
    grids = [
        (6.45, [("JA12 : PWR", "JA6 : PWR"), ("JA11 : GND", "JA5 : GND"), ("JA10 : G3", "JA4 : G2"),
                ("JA9 : H2", "JA3 : J2"), ("JA8 : K2", "JA2 : L2"), ("JA7 : H1", "JA1 : J1")]),
        (8.07, [("JXAC12 : PWR", "JXAC6 : PWR"), ("JXAC11 : GND", "JXAC5 : GND"),
                ("JXAC10 : N1", "JXAC4 : N2"), ("JXAC9 : M1", "JXAC3 : M2"),
                ("JXAC8 : M3", "JXAC2 : L3"), ("JXAC7 : K3", "JXAC1 : J3")]),
    ]
    hi = {"JA5 : GND": GRAY, "JA2 : L2": GREEN, "JA1 : J1": RED,
          "JXAC6 : PWR": GRAY, "JXAC5 : GND": GRAY, "JXAC3 : M2": ORANGE}

    def cell_color(label):
        if "PWR" in label:
            return "F4CCCC", RED
        if "GND" in label:
            return "D6E4F0", BLUE
        return "E7E6E6", "7F7F7F"

    for gy, rows in grids:
        for i, (l, r) in enumerate(rows):
            y = gy + i * pitch
            text_box(slide, lab_l, y, 1.42, pitch, [(l, 13, SUB, False)], align=PP_ALIGN.RIGHT, margin=0.04)
            for cx, lab in ((cell_x1, l), (cell_x2, r)):
                f, ln = cell_color(lab)
                if lab in hi:
                    f, ln = hi[lab], hi[lab]
                box(slide, cx, y + (pitch - cell) / 2, cell, cell, ln, fill=f,
                    shape=MSO_SHAPE.RECTANGLE, width=1)
            if r in hi:
                text_box(slide, lab_r, y + 0.01, lab_w, pitch - 0.02, [(r, 13, hi[r], True)],
                         shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=hi[r], fill="FFFFFF", margin=0.06)
                used[r] = (lab_r + lab_w, y + pitch / 2)
            else:
                text_box(slide, lab_r, y, lab_w, pitch, [(r, 13, SUB, False)], margin=0.06)
    # 사진 위 커넥터 표시
    for ry0, ry1 in ((0.24, 0.43), (0.52, 0.707)):
        x0, y0 = P(0.088, ry0, pic_h)
        x1, y1 = P(0.195, ry1, pic_h)
        box(slide, x0, y0, x1 - x0, y1 - y0, ON_BOARD[GRAY], width=2.5)

    # ---- 센서 콜아웃 (보드 위, 왼쪽 → 핀 라벨) ------------------------------
    sens = [
        ("JA5 : GND", GRAY, "JA5 · GND → SR04 GND", None),
        ("JA2 : L2", GREEN, "JA2 · L2 → SR04 ECHO (echo)", "echo 5 V 출력이면 분압 저항으로 3.3 V"),
        ("JA1 : J1", RED, "JA1 · J1 → SR04 TRIG (trig)", None),
        ("JXAC6 : PWR", GRAY, "JXADC6 · PWR (3.3 V) → DHT11 VCC", None),
        ("JXAC5 : GND", GRAY, "JXADC5 · GND → DHT11 GND", None),
        ("JXAC3 : M2", ORANGE, "JXADC3 · M2 → DHT11 DATA (dht11_io)", "PULLUP · 외부 5 kΩ 풀업 권장"),
    ]
    sx, sw, y = 4.75, 5.4, 2.62
    for key, color, title, sub in sens:
        h = 0.6 if sub else 0.38
        lines = [(title, 16, "FFFFFF", True)]
        if sub:
            lines.append((sub, 13, "FFFFFF", False))
        text_box(slide, sx, y, sw, h, lines, shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill=color, margin=0.14)
        tx, ty = used[key]
        wire(slide, sx, y + h / 2, tx, ty, color)
        y += h + 0.05

    # ---- 센서 전원 주의 (왼쪽 위) -------------------------------------------
    text_box(slide, 0.92, 2.62, 3.45, 2.3,
             [("⚠ 센서 전원", 17, RED, True),
              ("HC-SR04 는 보통 5 V 필요", 15, TEXT, False),
              ("(Pmod PWR 는 3.3 V)", 15, TEXT, False),
              ("→ 팀원 단독 테스트 때 쓴", 15, TEXT, False),
              ("    전원 연결 그대로 사용", 15, TEXT, False)],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill="FBEAEA", anchor=MSO_ANCHOR.MIDDLE, margin=0.2)

    # ---- 오른쪽 위 : 범례 ---------------------------------------------------
    # 한 항목 = 여러 줄 (줄바꿈 위치 직접 지정 → 단어 중간 끊김 방지)
    legend = [
        (BLUE, [[("sw0", 1), (" 시계 / 스톱워치 · ", 0), ("sw1", 1), (" 초.1/100초 / 시.분 · ", 0),
                 ("sw2", 1), (" 시계 시간 설정", 0)],
                [("sw3", 1), (" 12 / 24 시간 · ", 0), ("sw4", 1), (" 타이머", 0)]]),
        (TEAL, [[("sw5", 1), (" 거리 (SR04) 화면 · ", 0), ("sw6", 1), (" 온습도 (DHT11) 화면 (sw6 우선)", 0)],
                [("UART 로 /mod 를 받은 뒤에는 sw0 ~ sw6 무시", 0)]]),
        (ORANGE, [[("sw15", 1), (" 거리 막대 LED (LD0 ~ 7) 켜기", 0)]]),
        (TEAL, [[("버튼", 1), (" : 지금 화면의 모듈에만 들어감", 0)],
                [("(거리 · 온습도 : R 1 회 측정, U 자동 측정 켜기/끄기)", 0)]]),
    ]
    ly = 2.62
    for color, lines in legend:
        h = 0.3 * len(lines)
        box(slide, 10.6, ly + 0.06, 0.2, 0.2, color, fill=color, shape=MSO_SHAPE.RECTANGLE, width=0.75)
        text_box(slide, 10.95, ly, 8.15, h,
                 [[(t, 15, color if b else TEXT, bool(b)) for t, b in ln] for ln in lines],
                 anchor=MSO_ANCHOR.TOP)
        ly += h + 0.16

    # ---- USB ---------------------------------------------------------------
    ux, uy = P(0.323, 0.123, pic_h)
    x0, y0 = P(0.293, 0.085, pic_h)
    x1, y1 = P(0.355, 0.165, pic_h)
    box(slide, x0, y0, x1 - x0, y1 - y0, ON_BOARD[UART], width=3)
    cx, cy, cw, ch = 13.2, 5.25, 5.9, 0.64
    text_box(slide, cx, cy, cw, ch,
             [("micro USB : 전원 + UART (9600 8N1)", 16, "FFFFFF", True),
              ("rx = B18 (PC → FPGA) · tx = A18 (FPGA → PC)", 13, "FFFFFF", False)],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill=UART, margin=0.14)
    wire(slide, cx, cy + ch / 2, x1, (y0 + y1) / 2, UART)

    # ---- 버튼 ---------------------------------------------------------------
    btns = [
        ("U", 0.715, 0.525, TEAL, "BTNU · T18 → btn_U"),
        ("L", 0.663, 0.600, TEAL, "BTNL · W19 → btn_L"),
        ("C", 0.715, 0.603, RED, "BTNC · U18 → rst (리셋)"),
        ("R", 0.773, 0.600, TEAL, "BTNR · T17 → btn_R"),
        ("D", 0.715, 0.694, TEAL, "BTND · U17 → btn_D"),
    ]
    r = 0.17
    by, bh, bx, bw = 6.3, 0.48, 14.7, 4.4
    for i, (_, rx, ry, color, label) in enumerate(btns):
        px, py = P(rx, ry, pic_h)
        box(slide, px - r, py - r, 2 * r, 2 * r, ON_BOARD[color], shape=MSO_SHAPE.OVAL, width=3)
        yy = by + i * (bh + 0.1)
        text_box(slide, bx, yy, bw, bh, [(label, 16, "FFFFFF", True)],
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill=color, margin=0.14)
        wire(slide, bx, yy + bh / 2, px + r * 0.7, py, color)

    # ---- 보드 위 라벨 (FND · LED) --------------------------------------------
    fx, fy = P(0.218, 0.59, pic_h)
    text_box(slide, fx, fy - 0.36, 2.45, 0.3, [("FND (fnd_com · fnd_font)", 12, BLUE, True)],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill="FFFFFF", align=PP_ALIGN.CENTER)
    lx, ly2 = P(0.475, 0.742, pic_h)
    text_box(slide, lx, ly2 - 0.36, 1.9, 0.3, [("LED 0 ~ 15 (led)", 12, GREEN, True)],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill="FFFFFF", align=PP_ALIGN.CENTER)

    # ---- 스위치 -------------------------------------------------------------
    sw_color = {15: ORANGE, 6: TEAL, 5: TEAL, 4: BLUE, 3: BLUE, 2: BLUE, 1: BLUE, 0: BLUE}
    for k, color in sw_color.items():
        rx = 0.858 - k * 0.0451
        x0, y0 = P(rx - 0.021, 0.80, pic_h)
        x1, y1 = P(rx + 0.021, 0.955, pic_h)
        box(slide, x0, y0, x1 - x0, y1 - y0, ON_BOARD[color], width=3)
        pw = 0.0451 * PIC_W if k < 10 else 0.7
        text_box(slide, (x0 + x1) / 2 - pw / 2, y1 + 0.05, pw, 0.28, [(f"sw{k}", 13, color, True)],
                 align=PP_ALIGN.CENTER)

    # ---- 하단 주석 ----------------------------------------------------------
    text_box(slide, 12.6, 9.7, 6.5, 0.7,
             [("핀 이름 : merge_final_top.xdc 기준", 14, SUB, False),
              ("색 칸 = 이번 설계가 쓰는 핀 · 클럭 100 MHz (W5)", 14, SUB, False)],
             align=PP_ALIGN.RIGHT)


def keep(sh):
    """프레임 도형 · 보드 사진만 남김."""
    if sh.shape_type == 13:
        return True
    if sh.top < Inches(2.5):
        return True
    return sh.has_text_frame and re.fullmatch(r"\d\d", sh.text_frame.text.strip() or "")


def main():
    prs = Presentation(DECK)
    for slide in prs.slides:
        if any(sh.has_text_frame and sh.text_frame.text == TITLE and sh.top < Inches(2) for sh in slide.shapes):
            for sh in list(slide.shapes):
                if not keep(sh):
                    sh._element.getparent().remove(sh._element)
            build(slide)
            print("rebuilt", TITLE)
    prs.save(DECK)


if __name__ == "__main__":
    main()
