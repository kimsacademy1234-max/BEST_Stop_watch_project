# -*- coding: utf-8 -*-
"""'시스템 스펙' 슬라이드 — 시스템 경계도 형식 (PC ↔ FPGA ↔ 외부 센서 + 구현 결과 띠).

- 보드 IO 슬라이드 앞에 새 슬라이드로 넣고, 프레임 (머리글 · 섹션 · 제목 · 요약 · 쪽 번호) 은
  보드 IO 슬라이드에서 복사
- 다시 실행하면 본문만 다시 그림 (idempotent), 쪽 번호는 전체 다시 매김

    python3 spec_slide.py      (FINAL_presentation.pptx 제자리 수정)
"""
import copy
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from board_io_slide import DECK, GREEN, ORANGE, RED, SUB, TEXT, UART, box, text_box, wire
from board_io_slide import TITLE as BOARD_TITLE

TITLE = "시스템 스펙"
SECTION = "01  프로젝트 개요"
SUMMARY = "100 MHz FPGA 하나에 PC (UART) · 센서 2 종 연결 — LUT 7.2 % · WNS +0.320 ns 로 타이밍 위반 0"

DARK = "404040"
CARD = "F4F4F4"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ---------------------------------------------------------------- 도우미
def badge(slide, x, y, n, d=0.46, color=DARK):
    text_box(slide, x, y, d, d, [(str(n), 16, "FFFFFF", True)], shape=MSO_SHAPE.OVAL,
             fill=color, align=PP_ALIGN.CENTER)


def arrow2(slide, x1, y1, x2, y2, color, width=3):
    c = wire(slide, x1, y1, x2, y2, color, width)
    ln = c.line._get_or_add_ln()
    for tag in ("headEnd", "tailEnd"):
        el = ln.makeelement(A + tag, {"type": "triangle", "w": "med", "len": "med"})
        ln.append(el)
    return c


def set_first_run(shape, text):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def is_page_no(sh):
    return (sh.has_text_frame and re.fullmatch(r"\d\d", sh.text_frame.text.strip())
            and sh.top > Inches(10))


# ---------------------------------------------------------------- 본문
def build(s):
    # ---- FPGA 박스 + 내부 스펙 4 행 -----------------------------------------
    fx, fy, fw, fh = 6.75, 2.65, 6.9, 6.45
    box(s, fx, fy, fw, fh, DARK, width=2.5).adjustments[0] = 0.04
    text_box(s, fx + 0.25, fy + 0.1, fw - 0.5, 0.45,
             [[("FPGA", 18, TEXT, True), ("   Basys3 · Artix-7 XC7A35T", 14, SUB, False)]])
    rows = [
        (1, "클럭 · 리셋", ["W5 단일 클럭 · btnC 리셋", "reset_sync 하나로 동기 해제"], "100 MHz", "단일 클럭"),
        (3, "전송 주기 · 우선순위", ["전송 중 요청은 저장 후 전송", "우선순위 ERROR > get · auto"], "1 Hz", "자동 전송"),
        (4, "시간 해상도", ["디바운스 약 2 ms", "FND 한 바퀴 2 ms"], "1/100 s", "표시 단위"),
        (5, "센서 제어 · 오류 처리", ["DHT11 전원 후 1 s 대기", "실패 시 FND Err1~3 (원인별 status)"],
         "60 ms · 2 s", "SR04 · DHT11 주기"),
    ]
    ry, rh, gap = fy + 0.65, 1.3, 0.12
    row_y = []
    for n, title, detail, big, cap in rows:
        text_box(s, fx + 0.2, ry, fw - 0.4, rh, [("", 10, TEXT, False)],
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill=CARD)
        badge(s, fx + 0.35, ry + 0.14, n)
        text_box(s, fx + 0.95, ry + 0.12, 3.9, 0.45, [(title, 17, TEXT, True)])
        text_box(s, fx + 0.95, ry + 0.58, 3.95, 0.64,
                 [(t, 14, SUB, False) for t in detail], anchor=MSO_ANCHOR.TOP)
        text_box(s, fx + 4.6, ry + 0.2, 2.0, 0.55, [(big, 24, TEXT, True)], align=PP_ALIGN.RIGHT)
        text_box(s, fx + 4.6, ry + 0.78, 2.0, 0.35, [(cap, 13, SUB, False)], align=PP_ALIGN.RIGHT)
        row_y.append(ry + rh / 2)
        ry += rh + gap

    # ---- PC + UART 링크 -----------------------------------------------------
    px, py, pw, ph = 0.92, 3.3, 2.3, 5.3
    box(s, px, py, pw, ph, DARK, fill="FFFFFF", width=2.5).adjustments[0] = 0.06
    text_box(s, px, py + 0.2, pw, 0.9, [("PC", 24, TEXT, True), ("시리얼 터미널", 15, SUB, False)],
             align=PP_ALIGN.CENTER)
    # 터미널 모양 : UART 명령어 슬라이드의 사용 예시
    term = [("> /mod dist", "FFFFFF"), ("> /get", "FFFFFF"), ("> /btn U", "FFFFFF"),
            ("> /auto all", "FFFFFF"), ("> /set 24h", "FFFFFF")]
    tb = text_box(s, px + 0.18, py + 1.3, pw - 0.36, 3.7,
                  [(t, 16, c, True) for t, c in term], shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                  fill="262626", anchor=MSO_ANCHOR.TOP, margin=0.15)
    tb.adjustments[0] = 0.05
    tb.text_frame.margin_top = Inches(0.2)
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(22)
        for r in p.runs:
            r.font.name = "Consolas"
    ay = py + ph / 2
    arrow2(s, px + pw + 0.08, ay, fx - 0.08, ay, UART)
    lx, lw = px + pw + 0.2, fx - (px + pw) - 0.4
    badge(s, lx, ay - 1.75, 2, color=UART)
    text_box(s, lx + 0.6, ay - 1.75, lw - 0.6, 0.46, [("UART 인터페이스", 17, UART, True)])
    text_box(s, lx, ay - 1.1, lw, 0.6, [("9600 bps 8N1", 24, TEXT, True)], align=PP_ALIGN.CENTER)
    text_box(s, lx, ay + 0.25, lw, 1.2,
             [("16× 오버샘플링", 15, SUB, False),
              ("FIFO RX 8 / TX 64", 15, SUB, False),
              ("형식 오류 → ERROR", 15, SUB, False)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)

    # ---- 외부 센서 + 링크 ---------------------------------------------------
    sx, sy, sw, sh_ = 15.9, 3.3, 3.2, 5.3
    box(s, sx, sy, sw, sh_, DARK, fill="FFFFFF", width=2.5).adjustments[0] = 0.06
    badge(s, sx + 0.2, sy + 0.2, 6)
    text_box(s, sx + 0.8, sy + 0.15, sw - 0.9, 0.6,
             [("센서 성능", 17, TEXT, True), ("외부 센서 · 데이터시트", 12, SUB, False)])
    text_box(s, sx + 0.25, sy + 1.05, sw - 0.4, 1.0,
             [("SR04", 16, GREEN, True), ("2 ~ 400 cm", 24, TEXT, True)], anchor=MSO_ANCHOR.TOP)
    text_box(s, sx + 0.25, sy + 2.35, sw - 0.4, 1.45,
             [("DHT11", 16, ORANGE, True), ("±5 %RH", 24, TEXT, True),
              ("±2 ℃ (25 ℃ 기준)", 18, TEXT, True)], anchor=MSO_ANCHOR.TOP)
    text_box(s, sx + 0.25, sy + 4.4, sw - 0.4, 0.6,
             [[("실측 오차 ", 15, SUB, False), ("미측정", 15, RED, True)]])
    for y, color, label in ((sy + 1.55, GREEN, "trig · echo"), (sy + 3.05, ORANGE, "data")):
        arrow2(s, fx + fw + 0.08, y, sx - 0.08, y, color)
        text_box(s, fx + fw + 0.1, y - 0.42, sx - fx - fw - 0.2, 0.34, [(label, 14, color, True)],
                 align=PP_ALIGN.CENTER)

    # ---- 구현 결과 띠 -------------------------------------------------------
    by, bh = 9.3, 1.1
    text_box(s, 0.92, by, 18.18, bh, [("", 10, TEXT, False)], shape=MSO_SHAPE.ROUNDED_RECTANGLE, fill=CARD)
    badge(s, 1.15, by + 0.32, 7)
    text_box(s, 1.75, by + 0.15, 3.6, 0.8,
             [("FPGA 자원 · 타이밍", 17, TEXT, True), ("구현 결과 · Vivado 2022.2 · v2", 12, SUB, False)])
    stats = [
        ([("LUT ", 16, SUB, False), ("1,489", 24, TEXT, True), ("  7.2 %", 16, SUB, False)]),
        ([("FF ", 16, SUB, False), ("737", 24, TEXT, True), ("  1.8 %", 16, SUB, False)]),
        ([("WNS ", 16, SUB, False), ("+0.320 ns", 24, TEXT, True), ("  위반 0", 16, SUB, False)]),
        ([("전력 ", 16, SUB, False), ("미측정", 20, RED, True)]),
    ]
    x = 5.6
    for parts, w in zip(stats, (3.3, 2.8, 4.1, 2.6)):
        text_box(s, x, by, w, bh, [parts])
        x += w + 0.2


# ---------------------------------------------------------------- 슬라이드 준비
def find_slide(prs, title):
    for i, sl in enumerate(prs.slides):
        if any(sh.has_text_frame and sh.text_frame.text == title and sh.top < Inches(2) for sh in sl.shapes):
            return i, sl
    return None, None


def frame_of(slide):
    return [sh for sh in slide.shapes if (sh.top < Inches(2.5) and sh.shape_type != 13) or is_page_no(sh)]


def main():
    prs = Presentation(DECK)
    idx, slide = find_slide(prs, TITLE)
    if slide is None:
        bidx, board = find_slide(prs, BOARD_TITLE)
        slide = prs.slides.add_slide(board.slide_layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        for sh in frame_of(board):
            slide.shapes._spTree.append(copy.deepcopy(sh._element))
        # 보드 IO 앞으로 이동
        lst = prs.slides._sldIdLst
        new = lst[-1]
        lst.remove(new)
        lst.insert(bidx, new)
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text == BOARD_TITLE:
                set_first_run(sh, TITLE)
            elif sh.has_text_frame and sh.text_frame.text.startswith("PC 통신"):
                set_first_run(sh, SUMMARY)
    else:
        keep = [sh._element for sh in frame_of(slide)]
        for sh in list(slide.shapes):
            if not any(sh._element is k for k in keep):
                sh._element.getparent().remove(sh._element)
    build(slide)

    # 쪽 번호 다시 매김 (02 ~)
    for no, sl in enumerate(prs.slides, start=1):
        for sh in sl.shapes:
            if is_page_no(sh):
                set_first_run(sh, f"{no:02d}")
    prs.save(DECK)
    print("built", TITLE)


if __name__ == "__main__":
    main()
