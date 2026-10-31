# -*- coding: utf-8 -*-
"""사용자가 FINAL_presentation.pptx 에 직접 넣은 슬라이드를 덱 디자인으로 포맷.

- 프레임 : 우상단 머리글 · 섹션 표시 (● 01 프로젝트 개요) · Semilight 제목 ·
           [요약] 칩 + 한 줄 요약 · 우하단 두 자리 쪽 번호  (reference_ppt.pptx 3 쪽 형식)
- 색 : Google Slides 색 → 덱 팔레트, 카드 = F4F4F4 무테, 표 헤더 = 모듈색
- 폰트 : Courier New → Consolas (한글은 맑은 고딕)

    python3 format_slides.py      (FINAL_presentation.pptx 를 제자리 수정)
"""
import copy
import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "FINAL_presentation.pptx")
REF = os.path.join(HERE, "reference_ppt.pptx")

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

ORANGE = "C05621"   # 덱 팔레트 주황 (보드 IO 번호 · 강조)
UART = "6B46C1"     # UART 모듈색

# Google Slides 색 -> 덱 팔레트
COLOR_MAP = {
    "14182B": "262626",   # 본문
    "56606C": "595959",   # 보조 글
    "8A94A0": "8C8C8C",
    "F57C00": ORANGE,
    "F1F3F6": "F4F4F4",   # 카드
    "CDD3DA": "D9D9D9",   # 표 테두리
    "C62828": "C00000",
    "1F4E8C": UART,       # 표 헤더 = UART 모듈색
}

SLIDES = {
    # 제목 : (섹션, 요약, 슬라이드별 추가 조정 함수 이름)
    "Basys3 보드 IO 구성": (
        "01  프로젝트 개요",
        "PC 통신 1 · 센서 Pmod 2 · 버튼 5 · 스위치 8 · LED 16 · FND 4 자리 — 한 보드에 IO 배치",
        "fix_board_io",
    ),
    "UART 명령어 및 기능": (
        "01  프로젝트 개요",
        "명령 6 종 (/mod · /get · /btn · /auto · /set · /off) — PC 터미널로 화면 · 버튼 · 데이터 전송 제어",
        "fix_uart_cmd",
    ),
}


# ---------------------------------------------------------------- 공통
def inch(v):
    return Inches(v)


def recolor(el):
    for c in el.iter(A + "srgbClr"):
        v = c.get("val")
        if v in COLOR_MAP:
            c.set("val", COLOR_MAP[v])


def remove_border(sp, fill_hex="F4F4F4"):
    """덱 카드 형식 : 채움 F4F4F4 · 테두리 없음."""
    spPr = sp._element.spPr
    fills = [c for c in spPr.iter(A + "srgbClr") if c.get("val") == fill_hex]
    if not fills:
        return
    ln = spPr.find(A + "ln")
    if ln is not None:
        for ch in list(ln):
            ln.remove(ch)
        ln.append(ln.makeelement(A + "noFill", {}))


def fonts_mono(el):
    for rpr in el.iter(A + "rPr", A + "endParaRPr"):
        for tag in ("latin", "ea", "cs", "sym"):
            f = rpr.find(A + tag)
            if f is not None and f.get("typeface") == "Courier New":
                f.set("typeface", "Malgun Gothic" if tag == "ea" else "Consolas")


def set_text(shape, text, size=None):
    p = shape.text_frame.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    if size:
        p.runs[0].font.size = Pt(size)


def frame_shapes(ref):
    """reference 3 쪽에서 프레임 도형 원본 가져오기."""
    s = ref.slides[2]
    by = {sh.shape_id: sh for sh in s.shapes}
    return {
        "dot": by[3], "section": by[4], "title": by[5],
        "chip": by[6], "summary": by[7], "page": by[48],
    }


def add_clone(slide, src, x, y, w, h):
    el = copy.deepcopy(src._element)
    slide.shapes._spTree.append(el)
    sh = slide.shapes[-1]
    nid = max(s.shape_id for s in slide.shapes) + 1
    el.xpath("./*[1]/p:cNvPr")[0].set("id", str(nid))
    sh.left, sh.top, sh.width, sh.height = inch(x), inch(y), inch(w), inch(h)
    return sh


def apply_frame(slide, deck_header, fr, title, section, summary, page_no):
    # Google 제목 · 쪽 번호 placeholder 삭제
    for sh in list(slide.shapes):
        if sh.is_placeholder:
            sh._element.getparent().remove(sh._element)

    add_clone(slide, deck_header, 13.6, 0.22, 6.0, 0.3)          # 우상단 머리글 (14 pt)
    d = add_clone(slide, fr["dot"], 0.92, 0.6, 0.16, 0.16)
    s = add_clone(slide, fr["section"], 1.18, 0.46, 10.0, 0.42)
    set_text(s, section, 16)
    t = add_clone(slide, fr["title"], 0.92, 0.86, 18.2, 0.85)
    set_text(t, title, 40)
    c = add_clone(slide, fr["chip"], 0.92, 1.86, 1.2, 0.52)
    set_text(c, "요약", 18)
    m = add_clone(slide, fr["summary"], 2.3, 1.8, 16.9, 0.64)
    set_text(m, summary, 24)
    p = add_clone(slide, fr["page"], 18.2, 10.62, 0.9, 0.38)
    set_text(p, f"{page_no:02d}", 14)
    return d


def shift(slide, dy, skip=()):
    for sh in slide.shapes:
        if sh.shape_id in skip:
            continue
        sh.top = sh.top + inch(dy)


# ---------------------------------------------------------------- 슬라이드별
def fix_board_io(slide, content_ids):
    shift(slide, 0.62, skip=set(s.shape_id for s in slide.shapes) - content_ids)
    for sh in slide.shapes:
        if sh.shape_id in content_ids and sh.has_text_frame and "xdc" in sh.text_frame.text:
            for r in sh.text_frame.paragraphs[0].runs:
                r.font.size = Pt(16)


def fix_uart_cmd(slide, content_ids):
    tables = [sh for sh in slide.shapes if sh.has_table]
    tables.sort(key=lambda s: s.left)
    left, right = tables
    widths = {
        id(left): [1.0, 0.95, 2.4, 1.85, 1.12, 1.75],
        id(right): [1.0, 1.0, 2.15, 2.35, 1.17, 1.4],
    }
    for tb in (left, right):
        tb.top = inch(2.6)
        for col, w in zip(tb.table.columns, widths[id(tb)]):
            col.width = inch(w)
        for i, row in enumerate(tb.table.rows):
            if i and row.height < inch(0.5):
                row.height = inch(0.44)
    for sh in slide.shapes:
        if sh.shape_id not in content_ids or sh.has_table or not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if txt.startswith("※"):
            # 한 줄에 안 들어가 글자 단위로 끊기던 주석 -> 두 줄로 나눔
            sh.top, sh.height = inch(8.65), inch(0.8)
            tf = sh.text_frame
            full = "".join(r.text for r in tf.paragraphs[0].runs)
            a, b = full.split(" · /get", 1)
            p0 = tf.paragraphs[0]
            p1 = copy.deepcopy(p0._p)
            p0._p.addnext(p1)
            _replace_runs(p0._p, a)
            _replace_runs(p1, "※ /get" + b)
            for rpr in sh._element.iter(A + "rPr"):
                rpr.set("sz", "1800")
        elif txt.startswith("통신"):
            sh.top, sh.height = inch(9.5), inch(0.95)
    for sh in slide.shapes:
        if sh.shape_id in content_ids and sh.shape_type == 1 and sh.width > inch(15):
            sh.top, sh.height = inch(9.5), inch(0.95)         # 하단 통신 · 형식 카드


def _replace_runs(p, text):
    """문단 run 들을 원래 서식 (명령 = Consolas / 한글 = 맑은 고딕) 유지하며 text 로 교체.
    '/' 로 시작하는 토큰은 첫 Consolas run 서식, 나머지는 첫 맑은 고딕 run 서식."""
    runs = p.findall(A + "r")

    def face(r):
        f = r.find(A + "rPr/" + A + "latin")
        return f.get("typeface") if f is not None else ""

    mark = runs[0]                                    # '※' (주황)
    mono = next((r for r in runs if face(r) == "Consolas"), runs[0])
    norm = next((r for r in runs[1:] if face(r) != "Consolas" and "※" not in r.find(A + "t").text), runs[0])
    mark, mono, norm = copy.deepcopy(mark), copy.deepcopy(mono), copy.deepcopy(norm)
    for r in runs:
        p.remove(r)
    end = p.find(A + "endParaRPr")
    for tok in re.split(r"(※ ?|/[a-z]+)", text):
        if not tok:
            continue
        r = copy.deepcopy(mark if tok.startswith("※") else mono if tok.startswith("/") else norm)
        r.find(A + "t").text = tok
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)


# ---------------------------------------------------------------- main
def main():
    prs = Presentation(DECK)
    ref = Presentation(REF)
    fr = frame_shapes(ref)
    deck_header = next(sh for sh in prs.slides[1].shapes if sh.shape_id == 2)

    for no, slide in enumerate(prs.slides, start=1):
        title_ph = next((sh for sh in slide.shapes if sh.is_placeholder
                         and sh.has_text_frame and sh.text_frame.text in SLIDES), None)
        if title_ph is None:
            continue                       # 이미 포맷됨 or 대상 아님
        title = title_ph.text_frame.text
        section, summary, fn = SLIDES[title]

        content_ids = {sh.shape_id for sh in slide.shapes if not sh.is_placeholder}
        recolor(slide._element)
        fonts_mono(slide._element)
        for sh in slide.shapes:
            if sh.shape_type == 1:
                remove_border(sh)
        apply_frame(slide, deck_header, fr, title, section, summary, no)
        globals()[fn](slide, content_ids)
        print("formatted", no, title)

    prs.save(DECK)


if __name__ == "__main__":
    main()
