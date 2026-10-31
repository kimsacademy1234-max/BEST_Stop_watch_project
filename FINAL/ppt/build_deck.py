# -*- coding: utf-8 -*-
"""FINAL 발표 PPT 생성기.

디자인 베이스 = reference_ppt.pptx (v2.8 발표본) 의 표지 · 목차 슬라이드.
나머지 슬라이드는 지우고, 표지 · 목차 텍스트와 도형을 FINAL 내용으로 교체한다.
수정은 이 스크립트에서 하고 다시 실행 (pptx 직접 편집 X).

    python3 build_deck.py   ->  FINAL_presentation.pptx
"""
import copy
import os

from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "reference_ppt.pptx")
OUT = os.path.join(HERE, "FINAL_presentation.pptx")

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# ---------------------------------------------------------------- 내용
HEADER = "STOPWATCH · SENSOR · UART"

COVER_TITLE = ["STOPWATCH · SENSOR ·", "UART"]
COVER_SUB = "Basys-3 FPGA 시계 · 센서 · UART 통합 시스템 — 설계 · 검증"
COVER_DATE = "2026. 11. 02"
COVER_TEAM = "팀원   견다혜 · 이종혁 · 정보윤 · 정안성"

# 모듈 색 (표지 chip · 목차 범례 공통). 앞 3 색은 레퍼런스 그대로.
MODULES = [
    ("WTS",   "2B6CB0", "스톱워치 · 시계 · 타이머 (1/100 s)"),
    ("SR04",  "2F855A", "초음파 거리 0 ~ 400 cm · 60 ms 주기"),
    ("DHT11", "C05621", "온도 ℃ · 습도 %RH · 2 s 주기"),
    ("UART",  "6B46C1", "9600 bps · RX 8 / TX 64 FIFO"),
]

CONTENTS_SUMMARY = [
    "Basys-3 FPGA 한 보드에 시계 · 센서 · UART",
    "기능을 하나의 Top 으로 통합,",
    "Top · 하위 모듈 테스트벤치로 전체 경로 검증",
]

CONTENTS = [
    ("01", "프로젝트 개요", "요약 · 목적 · 스펙 · 보드 IO · UART 명령표"),
    ("02", "시스템 구조", "Top 블록 다이어그램 · 신호 흐름"),
    ("03", "Top 검증", "다이어그램 흐름 · 전체 동작 파형 · Edge case · TB 신뢰성"),
    ("04", "하위 모듈 검증", "모듈별 다이어그램 · Edge case · TB 결과"),
    ("05", "트러블 슈팅", "현상 → 원인 → 해결 → 적용 결과"),
    ("06", "시연 영상", "보드 동작"),
    ("07", "결론 · 고찰", "목표 달성 수치 · 남은 과제"),
]


# ---------------------------------------------------------------- 도우미
def keep_slides(prs, keep):
    """keep (0-based index) 이외 슬라이드 삭제. 저장 시 고아 파트는 빠짐."""
    sld_ids = prs.slides._sldIdLst
    for i, sid in reversed(list(enumerate(list(sld_ids)))):
        if i not in keep:
            prs.part.drop_rel(sid.rId)
            sld_ids.remove(sid)


def by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def set_lines(shape, lines):
    """문단 서식 유지하며 텍스트 교체 (첫 run 서식 사용)."""
    txb = shape.text_frame._txBody
    paras = txb.findall(A + "p")
    while len(paras) < len(lines):
        new = copy.deepcopy(paras[-1])
        paras[-1].addnext(new)
        paras.append(new)
    for p, text in zip(paras, lines):
        runs = p.findall(A + "r")
        runs[0].find(A + "t").text = text
        for r in runs[1:]:
            p.remove(r)
    for p in paras[len(lines):]:
        txb.remove(p)


def set_size(shape, pt):
    """모든 run · 문단 끝 글자 크기 변경."""
    for tag in ("rPr", "endParaRPr"):
        for el in shape._element.iter(A + tag):
            el.set("sz", str(int(pt * 100)))


def resize(shape, w=None, h=None):
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)


def set_color(shape, hex_):
    for el in shape._element.iter(A + "srgbClr"):
        el.set("val", hex_)


def clone(slide, src, x=None, y=None):
    el = copy.deepcopy(src._element)
    slide.shapes._spTree.append(el)
    sh = slide.shapes[-1]
    # shape id 중복 방지
    el.xpath("./*[1]/p:cNvPr")[0].set("id", str(max(s.shape_id for s in slide.shapes) + 1))
    if x is not None:
        sh.left = Inches(x)
    if y is not None:
        sh.top = Inches(y)
    return sh


def place(shape, x=None, y=None):
    if x is not None:
        shape.left = Inches(x)
    if y is not None:
        shape.top = Inches(y)


# ---------------------------------------------------------------- 표지
def build_cover(s):
    set_lines(by_id(s, 2), [HEADER])
    set_lines(by_id(s, 3), COVER_TITLE)
    set_lines(by_id(s, 4), [COVER_SUB])
    set_lines(by_id(s, 8), [COVER_DATE])
    set_lines(by_id(s, 9), [COVER_TEAM])

    # 글자 크기 (20 x 11.25 in 판 = 일반 16:9 의 1.5 배 → 크기도 1.5 배 기준)
    set_size(by_id(s, 2), 14)
    title = by_id(s, 3)
    set_size(title, 72)
    place(title, y=2.2)
    resize(title, w=17.0, h=2.9)
    sub = by_id(s, 4)
    set_size(sub, 28)
    place(sub, y=5.25)
    resize(sub, w=17.0, h=0.7)
    for sid, y in ((8, 8.75), (9, 9.3)):
        sh = by_id(s, sid)
        set_size(sh, 22)
        place(sh, y=y)
        resize(sh, h=0.5)

    chips = [by_id(s, 5), by_id(s, 6), by_id(s, 7)]
    chips.append(clone(s, chips[0]))
    for i, (chip, (name, color, _)) in enumerate(zip(chips, MODULES)):
        place(chip, x=0.95 + i * 2.6, y=6.3)
        resize(chip, w=2.35, h=0.62)
        set_size(chip, 20)
        set_lines(chip, [name])
        set_color(chip, color)


# ---------------------------------------------------------------- 목차
def build_contents(s):
    set_lines(by_id(s, 2), [HEADER])
    set_lines(by_id(s, 4), CONTENTS_SUMMARY)
    set_size(by_id(s, 2), 14)
    set_size(by_id(s, 3), 56)
    summ = by_id(s, 4)
    set_size(summ, 20)
    place(summ, y=2.45)
    resize(summ, w=7.3, h=1.5)
    for el in summ._element.iter(A + "lnSpc"):
        el[0].set("val", "125000")
    set_size(by_id(s, 42), 14)

    # 레퍼런스 6 행 (번호, 제목, 부제, 구분선) -> 7 행으로 재배치
    rows = [tuple(by_id(s, b + k) for k in range(4)) for b in (5, 9, 13, 17, 21, 25)]
    rows.append(tuple(clone(s, sh) for sh in rows[-1]))
    y0, pitch = 1.30, 1.24
    for i, ((num, title, sub, line), (n, t, d)) in enumerate(zip(rows, CONTENTS)):
        y = y0 + i * pitch
        place(num, y=y)
        place(title, y=y - 0.02)
        place(sub, y=y + 0.6)
        place(line, y=y + 1.02)
        set_size(num, 40)
        set_size(title, 30)
        set_size(sub, 18)
        resize(title, w=9.0, h=0.55)
        resize(sub, w=9.1, h=0.4)
        set_lines(num, [n])
        set_lines(title, [t])
        set_lines(sub, [d])

    # 왼쪽 모듈 범례 : 레퍼런스 3 개 (점, 이름, 설명) -> 4 개
    legend = [tuple(by_id(s, b + k) for k in range(3)) for b in (36, 33, 39)]
    legend.append(tuple(clone(s, sh) for sh in legend[-1]))
    ly0, lpitch = 5.3, 1.15
    for i, ((dot, name, desc), (m, color, text)) in enumerate(zip(legend, MODULES)):
        y = ly0 + i * lpitch
        place(dot, y=y + 0.13)
        place(name, y=y)
        place(desc, y=y + 0.5)
        set_size(name, 22)
        set_size(desc, 18)
        resize(name, h=0.5)
        resize(desc, w=6.9, h=0.42)
        dot.width = dot.height = Inches(0.26)
        set_lines(name, [m])
        set_lines(desc, [text])
        set_color(dot, color)
        set_color(name, color)


def main():
    # 사용자가 직접 슬라이드를 추가한 뒤에는 덮어쓰지 않음
    if os.path.exists(OUT) and len(Presentation(OUT).slides) > 2:
        raise SystemExit(f"{OUT} 에 사용자 슬라이드가 있음 — 덮어쓰기 중단")
    prs = Presentation(SRC)
    keep_slides(prs, keep={0, 1})
    build_cover(prs.slides[0])
    build_contents(prs.slides[1])
    prs.save(OUT)
    print("saved", OUT, f"{os.path.getsize(OUT) / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
